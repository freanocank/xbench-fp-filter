import io
import unittest
from unittest.mock import patch

import openpyxl

from ui.backend import (
    RunConfig, RowResult, build_filtered_xlsx, start_processing_task,
    test_api_connection as check_api_connection,
)


class BackendTests(unittest.TestCase):
    def test_connection_check_has_a_deadline_and_closes_its_client(self):
        cfg = RunConfig(api_key="local-test-key", api_base="http://127.0.0.1:9999/v1")
        with patch("openai.OpenAI") as factory:
            client = factory.return_value.__enter__.return_value
            client.chat.completions.create.return_value.model = "test-model"
            ok, message = check_api_connection(cfg)
        self.assertTrue(ok)
        self.assertIn("test-model", message)
        factory.assert_called_once_with(
            api_key=cfg.api_key, base_url=cfg.api_base, timeout=15, max_retries=0,
        )
        factory.return_value.__exit__.assert_called_once()

    def test_failed_file_finishes_task_and_does_not_block_next_file(self):
        successful = {"filename": "second.xlsx", "rows": [], "stats": {"total": 0}}
        with patch("ui.backend.process_file", side_effect=[RuntimeError("bad endpoint"), successful]):
            task = start_processing_task([("first.xlsx", b""), ("second.xlsx", b"")], RunConfig())
            self.assertTrue(task._done_event.wait(3), "Background worker did not finish")
        self.assertEqual(len(task.results), 2)
        self.assertIn("bad endpoint", task.results[0]["error"])
        self.assertEqual(task.results[1], successful)

    def test_export_retains_unknown_verdicts_and_removes_only_false_positives(self):
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet["A11"] = "Tag Mismatch"
        for index in (12, 13, 14):
            sheet.cell(index, 1, f"sample.txt ({index})")
            sheet.cell(index, 3, f"source-{index}")
        source = io.BytesIO()
        workbook.save(source)
        workbook.close()
        rows = [RowResult(n, "Tag Mismatch", "tags", "", "", verdict, "")
                for n, verdict in [(12, "误报"), (13, "真错误"), (14, None)]]
        exported = openpyxl.load_workbook(io.BytesIO(build_filtered_xlsx(source.getvalue(), rows)))
        values = [exported.active.cell(n, 3).value for n in range(11, exported.active.max_row + 1)]
        self.assertNotIn("source-12", values)
        self.assertIn("source-13", values)
        self.assertIn("source-14", values)
        exported.close()
