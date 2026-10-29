import base64
import io
import unittest
import zipfile
import zlib
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock, patch

import openpyxl
from openpyxl.cell.rich_text import CellRichText, TextBlock

from core.excel_compat import normalize_report_bytes
from ui.backend import RunConfig, build_filtered_xlsx, build_zip_bytes, process_file


def legacy_report():
    # Generated with xlwt; contains only synthetic Chinese/Korean report text.
    path = Path(__file__).parent / "fixtures" / "legacy_report.xls.zlib.b64"
    return zlib.decompress(base64.b64decode(path.read_bytes()))


def theme_zip():
    """Small synthetic theme payload; contains no workbook or customer data."""
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("[Content_Types].xml", '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"/>')
        archive.writestr("theme/theme/theme1.xml", "<theme/>")
    # Include the completed central directory when computing OLE alignment.
    size = len(buffer.getvalue())
    with zipfile.ZipFile(buffer, "a") as archive:
        archive.comment = b" " * (-size % 512)
    return buffer.getvalue()


class ExcelCompatibilityTests(unittest.TestCase):
    def test_xls_with_theme_zip_is_read_as_xls_not_as_an_xlsx_archive(self):
        content = legacy_report() + theme_zip()
        self.assertEqual(len(content) % 512, 0)
        self.assertTrue(zipfile.is_zipfile(io.BytesIO(content)))
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            self.assertNotIn("xl/workbook.xml", archive.namelist())
        normalized = normalize_report_bytes(content)
        self.assertIsNot(normalized, content)
        workbook = openpyxl.load_workbook(io.BytesIO(normalized), rich_text=True)
        self.assertEqual(str(workbook["Xbench QA"]["C13"].value), "保留红色标记文本 한국어")
        self.assertEqual(workbook["Xbench QA"].row_dimensions[13].height, 88)
        workbook.close()

    def test_theme_only_zip_has_an_actionable_workbook_error(self):
        with self.assertRaisesRegex(ValueError, "未找到有效的 Excel 工作簿"):
            normalize_report_bytes(theme_zip())

    def test_xls_conversion_preserves_text_highlights_layout_and_sheet_names(self):
        normalized = normalize_report_bytes(legacy_report())
        workbook = openpyxl.load_workbook(io.BytesIO(normalized), rich_text=True)
        self.assertEqual(workbook.sheetnames, ["Xbench QA", "说明"])
        sheet = workbook["Xbench QA"]
        value = sheet["C13"].value
        self.assertIsInstance(value, CellRichText)
        self.assertEqual(str(value), "保留红色标记文本 한국어")
        red_runs = [run.text for run in value if isinstance(run, TextBlock)
                    and run.font.color is not None and run.font.color.rgb == "FFFF0000"]
        self.assertEqual(red_runs, ["红色标记"])
        self.assertEqual(sheet.row_dimensions[13].height, 88)
        self.assertEqual(sheet.column_dimensions["C"].width, 42)
        self.assertIn("A1:F1", sheet.merged_cells)
        self.assertEqual(sheet.freeze_panes, "A12")
        self.assertTrue(sheet["C13"].alignment.wrap_text)
        self.assertEqual(sheet["C13"].border.bottom.style, "thin")
        self.assertEqual(sheet["C14"].value, "=literal localization text")
        self.assertEqual(sheet["C14"].data_type, "s")
        workbook.close()

    def test_native_xlsx_bytes_are_unchanged_even_when_named_xls(self):
        workbook = openpyxl.Workbook()
        buffer = io.BytesIO()
        workbook.save(buffer)
        workbook.close()
        original = buffer.getvalue()
        self.assertIs(normalize_report_bytes(original), original)

    def process_legacy_report(self):
        async def review(**kwargs):
            label = "误报" if kwargs["source_text"] == "验收样例：紫色花卉" else "真错误"
            return SimpleNamespace(final_label=label, llm_analysis="合成测试判定")

        with patch("ui.backend.apply_config"), patch(
            "core.llm_review.llm_secondary_review", side_effect=review,
        ) as mock_review:
            result = process_file(legacy_report(), "旧版报告.xls", RunConfig())
        self.assertNotIn("error", result)
        self.assertEqual(mock_review.call_count, 3)
        self.assertEqual([row.row_num for row in result["rows"]], [12, 13, 14])
        self.assertEqual(result["stats"], {"total": 3, "suppressed": 1, "kept": 2})
        return result

    def test_xls_is_parsed_and_reviewed_before_filtered_xlsx_export(self):
        result = self.process_legacy_report()
        content = build_filtered_xlsx(legacy_report(), result["rows"])
        workbook = openpyxl.load_workbook(io.BytesIO(content), rich_text=True)
        sheet = workbook["Xbench QA"]
        self.assertEqual(sheet["A12"].value, "sample.txt (2)")
        self.assertEqual(sheet["A13"].value, "sample.txt (3)")
        self.assertEqual(str(sheet["C12"].value), "保留红色标记文本 한국어")
        self.assertIsInstance(sheet["C12"].value, CellRichText)
        self.assertEqual(sheet.row_dimensions[12].height, 88)
        workbook.close()

    def test_xls_results_are_included_in_zip_as_real_xlsx_files(self):
        result = self.process_legacy_report()
        content = build_zip_bytes([result], {"旧版报告.xls": legacy_report()})
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            self.assertEqual(set(archive.namelist()), {"旧版报告_filtered.xlsx", "旧版报告_analysis.xlsx"})
            for name in archive.namelist():
                workbook = openpyxl.load_workbook(io.BytesIO(archive.read(name)), rich_text=True)
                workbook.close()

    def test_invalid_xls_fails_before_applying_api_configuration(self):
        with patch("ui.backend.apply_config") as apply, patch(
            "core.llm_review.llm_secondary_review", new_callable=AsyncMock,
        ) as review:
            result = process_file(b"not an excel file", "broken.xls", RunConfig())
        self.assertIn("解析失败", result["error"])
        self.assertIn("另存为 .xlsx", result["error"])
        apply.assert_not_called()
        review.assert_not_awaited()

    def test_empty_report_has_an_actionable_error(self):
        with self.assertRaisesRegex(ValueError, "报告文件为空"):
            normalize_report_bytes(b"")
