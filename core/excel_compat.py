"""Normalize legacy Excel QA reports to XLSX without dropping rich-text runs.

Native XLSX bytes pass through unchanged. XLS conversion reads BIFF data and
common report formatting; it does not require Excel or LibreOffice at runtime.
"""

import io
import zipfile

import openpyxl
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.utils import get_column_letter

_OLE_SIGNATURE = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
_UNDERLINES = {1: "single", 2: "double", 33: "singleAccounting", 34: "doubleAccounting"}
_BORDERS = {
    1: "thin", 2: "medium", 3: "dashed", 4: "dotted", 5: "thick", 6: "double",
    7: "hair", 8: "mediumDashed", 9: "dashDot", 10: "mediumDashDot",
    11: "dashDotDot", 12: "mediumDashDotDot", 13: "slantDashDot",
}
_PATTERNS = {
    1: "solid", 2: "mediumGray", 3: "darkGray", 4: "lightGray",
    5: "darkHorizontal", 6: "darkVertical", 7: "darkDown", 8: "darkUp",
    9: "darkGrid", 10: "darkTrellis", 11: "lightHorizontal", 12: "lightVertical",
    13: "lightDown", 14: "lightUp", 15: "lightGrid", 16: "lightTrellis",
    17: "gray125", 18: "gray0625",
}


def _color(book, index):
    rgb = book.colour_map.get(index)
    return "FF" + "".join(f"{value:02X}" for value in rgb) if rgb else None


def _font_options(book, font):
    return {
        "sz": font.height / 20,
        "b": bool(font.bold), "i": bool(font.italic),
        "strike": bool(font.struck_out),
        "outline": bool(font.outline), "shadow": bool(font.shadow),
        "u": _UNDERLINES.get(font.underline_type),
        "color": _color(book, font.colour_index),
        "vertAlign": {1: "superscript", 2: "subscript"}.get(font.escapement),
    }


def _cell_style(book, xf):
    font = book.font_list[xf.font_index]
    align = xf.alignment
    border = xf.border

    def side(name):
        return Side(
            style=_BORDERS.get(getattr(border, f"{name}_line_style")),
            color=_color(book, getattr(border, f"{name}_colour_index")),
        )

    return {
        "font": Font(name=font.name, **_font_options(book, font)),
        "alignment": Alignment(
            horizontal={0: "general", 1: "left", 2: "center", 3: "right", 4: "fill",
                        5: "justify", 6: "centerContinuous", 7: "distributed"}.get(align.hor_align),
            vertical={0: "top", 1: "center", 2: "bottom", 3: "justify", 4: "distributed"}.get(align.vert_align),
            textRotation=align.rotation, wrap_text=bool(align.text_wrapped),
            shrink_to_fit=bool(align.shrink_to_fit), indent=align.indent_level,
        ),
        "border": Border(
            left=side("left"), right=side("right"), top=side("top"), bottom=side("bottom"),
            diagonal=side("diag"), diagonalUp=bool(border.diag_up), diagonalDown=bool(border.diag_down),
        ),
        "fill": PatternFill(
            patternType=_PATTERNS.get(xf.background.fill_pattern),
            fgColor=_color(book, xf.background.pattern_colour_index) or "00000000",
            bgColor=_color(book, xf.background.background_colour_index) or "00000000",
        ),
        "protection": Protection(locked=bool(xf.protection.cell_locked), hidden=bool(xf.protection.formula_hidden)),
        "number_format": book.format_map[xf.format_key].format_str if xf.format_key in book.format_map else "General",
    }


def _rich_text(book, text, runs):
    parts = []
    if runs[0][0]:
        parts.append(text[:runs[0][0]])
    for index, (start, font_index) in enumerate(runs):
        stop = runs[index + 1][0] if index + 1 < len(runs) else len(text)
        if stop > start:
            font = book.font_list[font_index]
            parts.append(TextBlock(InlineFont(rFont=font.name, **_font_options(book, font)), text[start:stop]))
    return CellRichText(parts)


def _copy_sheet(book, source, target, styles):
    import xlrd

    for row_index in range(source.nrows):
        for col_index in range(source.row_len(row_index)):
            cell = source.cell(row_index, col_index)
            if cell.ctype == xlrd.XL_CELL_EMPTY:
                continue
            target_cell = target.cell(row_index + 1, col_index + 1)
            value = cell.value
            if cell.ctype == xlrd.XL_CELL_DATE:
                value = xlrd.xldate_as_datetime(value, book.datemode)
            elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                value = bool(value)
            elif cell.ctype == xlrd.XL_CELL_ERROR:
                value = xlrd.error_text_from_code.get(value, "#VALUE!")
            runs = source.rich_text_runlist_map.get((row_index, col_index))
            if runs and cell.ctype == xlrd.XL_CELL_TEXT:
                value = _rich_text(book, value, runs)
            target_cell.value = value
            if cell.ctype == xlrd.XL_CELL_TEXT:
                # Report text beginning with '=' is text, not a new formula.
                target_cell.data_type = "s"
            elif cell.ctype == xlrd.XL_CELL_ERROR:
                target_cell.data_type = "e"
            if cell.xf_index not in styles:
                styles[cell.xf_index] = _cell_style(book, book.xf_list[cell.xf_index])
            for attribute, setting in styles[cell.xf_index].items():
                setattr(target_cell, attribute, setting)

    for row_index, info in source.rowinfo_map.items():
        dimension = target.row_dimensions[row_index + 1]
        dimension.height = info.height / 20
        dimension.hidden = bool(info.hidden)
        dimension.outlineLevel = info.outline_level
    for col_index, info in source.colinfo_map.items():
        dimension = target.column_dimensions[get_column_letter(col_index + 1)]
        dimension.width = info.width / 256
        dimension.hidden = bool(info.hidden)
        dimension.outlineLevel = info.outline_level
    for row_start, row_stop, col_start, col_stop in source.merged_cells:
        target.merge_cells(start_row=row_start + 1, end_row=row_stop,
                           start_column=col_start + 1, end_column=col_stop)
    target.sheet_state = {0: "visible", 1: "hidden", 2: "veryHidden"}.get(source.visibility, "visible")
    if source.panes_are_frozen:
        target.freeze_panes = target.cell(source.horz_split_pos + 1, source.vert_split_pos + 1)


def normalize_report_bytes(content: bytes) -> bytes:
    """Accept XLS/XLSX content; return XLSX bytes with original row positions."""
    if not content:
        raise ValueError("报告文件为空，请重新导出后上传。")
    # XLS theme streams may themselves contain a valid ZIP archive. Inspect
    # the outer OLE container first; is_zipfile alone also finds those streams.
    if not content.startswith(_OLE_SIGNATURE) and zipfile.is_zipfile(io.BytesIO(content)):
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if "xl/workbook.xml" not in archive.namelist():
                raise ValueError("文件中未找到有效的 Excel 工作簿，请用 Excel 或 WPS 另存为 .xlsx 后重试。")
        return content

    import xlrd
    try:
        book = xlrd.open_workbook(file_contents=content, formatting_info=True, on_demand=True)
    except xlrd.XLRDError as exc:
        raise ValueError("无法读取该 Excel 报告，请用 Excel 或 WPS 另存为 .xlsx 后重试。") from exc
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    try:
        styles = {}
        for source in book.sheets():
            _copy_sheet(book, source, workbook.create_sheet(source.name), styles)
        buffer = io.BytesIO()
        workbook.save(buffer)
        return buffer.getvalue()
    finally:
        book.release_resources()
        workbook.close()
